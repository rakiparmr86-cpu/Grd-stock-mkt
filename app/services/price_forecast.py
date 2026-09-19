"""Price history + statistical forecast for one ticker, and its HTML / Excel exports.

The forecast is the calculation engine's own ETS/ARIMA (``statistics.forecast``)
over daily closes: a trend extrapolation with a 95% interval, not a market
prediction. No LLM.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime
from typing import Any

import pandas as pd

from app.services.calculations.statistics import forecast as stat_forecast
from app.services.market_data.repository import load_ohlcv_frame
from app.services.reports.charts import _fig_to_b64
from app.services.reports.renderer import ReportRenderer

MIN_BARS = 30


class ForecastError(ValueError):
    pass


def build_price_forecast(ticker: str, *, history: int = 120, ahead: int = 20) -> dict[str, Any]:
    frame = load_ohlcv_frame(ticker.upper(), limit=max(history, 250))
    if frame.empty:
        raise ForecastError(f"no price data for {ticker.upper()}")
    if len(frame) < MIN_BARS:
        raise ForecastError(
            f"{ticker.upper()} has {len(frame)} bars; at least {MIN_BARS} are needed to forecast"
        )
    closes = frame["close"].astype(float)
    fc = stat_forecast(closes, periods_ahead=ahead, method="auto")
    if fc.get("insufficient_data"):
        raise ForecastError(fc.get("reason", "not enough data"))

    last_ts = pd.Timestamp(frame.index[-1])
    dates = pd.bdate_range(last_ts + pd.Timedelta(days=1), periods=ahead)
    points = [
        {"ts": d.isoformat(), "value": v, "low": ci["low"], "high": ci["high"]}
        for d, v, ci in zip(dates, fc["forecast"], fc["confidence_intervals"], strict=True)
    ]
    tail = frame.tail(history)
    bars = [
        {"ts": pd.Timestamp(ts).isoformat(), "open": float(r.open), "high": float(r.high),
         "low": float(r.low), "close": float(r.close), "volume": float(r.volume)}
        for ts, r in tail.iterrows()
    ]
    last = float(closes.iloc[-1])
    end = points[-1]["value"]
    return {
        "ticker": ticker.upper(),
        "method": fc["method"],
        "ahead": ahead,
        "last_close": last,
        "last_ts": last_ts.isoformat(),
        "forecast_end": end,
        "forecast_change_pct": (end - last) / last * 100 if last else None,
        "residual_std": fc["residual_std"],
        "n_observations": fc["n_observations"],
        "bars": bars,
        "forecast": points,
        "note": ("Statistical trend extrapolation (Holt-Winters / ARIMA) with a 95% interval. "
                 "The band widens because uncertainty grows with distance; it is not advice."),
    }


def _chart_b64(fc: dict[str, Any]) -> str:
    import matplotlib.pyplot as plt

    hist = pd.DataFrame(fc["bars"])
    hist["ts"] = pd.to_datetime(hist["ts"])
    proj = pd.DataFrame(fc["forecast"])
    proj["ts"] = pd.to_datetime(proj["ts"])
    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.plot(hist["ts"], hist["close"], color="#2563eb", linewidth=1.6, label="close")
    ax.plot(
        [hist["ts"].iloc[-1], *proj["ts"]], [hist["close"].iloc[-1], *proj["value"]],
        color="#d97706", linestyle="--", linewidth=1.6, label=f"forecast ({fc['method']})",
    )
    ax.fill_between(proj["ts"], proj["low"], proj["high"], color="#d97706", alpha=0.18,
                    label="95% interval")
    ax.set_title(f"{fc['ticker']} — price and {fc['ahead']}-day forecast")
    ax.grid(alpha=0.25)
    ax.legend(loc="upper left", fontsize=8)
    fig.autofmt_xdate()
    return _fig_to_b64(fig)


def forecast_html(fc: dict[str, Any]) -> str:
    def n(v: float) -> str:
        return f"{v:,.2f}"

    step = max(1, len(fc["forecast"]) // 10)
    rows = [[p["ts"][:10], n(p["value"]), n(p["low"]), n(p["high"])]
            for p in fc["forecast"][::step]]
    if fc["forecast"][-1]["ts"][:10] != rows[-1][0]:
        p = fc["forecast"][-1]
        rows.append([p["ts"][:10], n(p["value"]), n(p["low"]), n(p["high"])])
    change = fc["forecast_change_pct"]
    payload = {
        "title": f"{fc['ticker']} — price forecast",
        "run_id": "—",
        "chart_b64": _chart_b64(fc),
        "tables": [
            {"heading": "Summary", "note": fc["note"], "num": True,
             "columns": ["Metric", "Value"],
             "rows": [
                 ["Last close", n(fc["last_close"])],
                 [f"Forecast in {fc['ahead']} bars", n(fc["forecast_end"])],
                 ["Change vs last close", "—" if change is None else f"{change:+.1f}%"],
                 ["Method", fc["method"]],
                 ["Observations used", str(fc["n_observations"])],
             ]},
            {"heading": "Forecast path", "note": "", "num": True,
             "columns": ["Date", "Forecast", "95% low", "95% high"], "rows": rows},
        ],
        "sections": [],
    }
    return ReportRenderer().render_html(payload)


def forecast_excel(fc: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Forecast"
    head = PatternFill("solid", fgColor="D9E2F3")
    ws.append([f"{fc['ticker']} — price and forecast"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([fc["note"]])
    ws.append([])
    ws.append(["Date", "Type", "Actual close", "Forecast", "95% low", "95% high"])
    for c in ws[4]:
        c.font, c.fill = Font(bold=True), head
    for i, b in enumerate(fc["bars"]):
        is_last = i == len(fc["bars"]) - 1
        # repeat the last actual in the forecast column so the two lines join
        ws.append([b["ts"][:10], "actual", b["close"], b["close"] if is_last else None, None, None])
    for p in fc["forecast"]:
        ws.append([p["ts"][:10], "forecast", None, p["value"], p["low"], p["high"]])
    first, last = 5, ws.max_row
    for row in ws.iter_rows(min_row=first, max_row=last, min_col=3, max_col=6):
        for c in row:
            c.number_format = "#,##0.00"
    for col, w in zip("ABCDEF", [14, 10, 14, 12, 12, 12], strict=True):
        ws.column_dimensions[col].width = w

    chart = LineChart()
    chart.title = f"{fc['ticker']} close and forecast"
    chart.height, chart.width = 9, 22
    chart.y_axis.title, chart.x_axis.title = "Price", "Date"
    chart.add_data(
        Reference(ws, min_col=3, max_col=6, min_row=4, max_row=last), titles_from_data=True,
    )
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    styles = [("2563EB", None), ("D97706", "dash"), ("F2B36B", "sysDot"), ("F2B36B", "sysDot")]
    for series, (color, dash) in zip(chart.series, styles, strict=True):
        series.graphicalProperties.line.solidFill = color
        series.smooth = False
        if dash:
            series.graphicalProperties.line.dashStyle = dash
    ws.add_chart(chart, "H4")

    summary = wb.create_sheet("Summary")
    for k, v in [
        ("Ticker", fc["ticker"]), ("Method", fc["method"]),
        ("Last close", fc["last_close"]), ("Last bar date", fc["last_ts"][:10]),
        (f"Forecast in {fc['ahead']} bars", fc["forecast_end"]),
        ("Change vs last close %", fc["forecast_change_pct"]),
        ("Observations used", fc["n_observations"]),
        ("Generated", datetime.now(UTC).strftime("%Y-%m-%d %H:%M UTC")),
    ]:
        summary.append([k, v])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
