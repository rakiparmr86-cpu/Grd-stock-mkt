# ruff: noqa: E501, RUF001
"""Prediction Inputs + Prediction Report for a Screener-style workbook.

Reproduces the layout of the hand-built ``GRDPrediction_Model.xlsx`` sample
(inputs sheet, momentum snapshot, Bear/Base/Bull FY+1 scenarios, multi-year
outlook, confidence) but computes it from any Screener export — no LLM. Scenario
assumptions are derived from the company's own history and written as editable
cells, so changing one in Excel recalculates the report sheet.
"""

from __future__ import annotations

import math
from pathlib import Path
from typing import Any

import pandas as pd

from app.services.calculations.scenario import (
    fiscal_year_labels,
    multi_year_outlook,
    scenario_projection,
)
from app.services.inputs.screener_excel import (
    parse_label_value_block,
    read_statement_sheet_raw,
    transpose_statement_sheet,
)

REQUIRED_SHEETS = {"Profit & Loss", "Quarters", "Data Sheet"}
EXTERNAL_PARAMS = [
    ("RBI Repo Rate", "Funding cost / margin environment"),
    ("GDP Growth Forecast", "Credit demand and asset quality"),
    ("Loan / Business Growth", "Core revenue growth driver"),
    ("Deposit / Funding Growth", "Funding capacity and liquidity"),
    ("Net Interest Margin (NIM)", "Profitability / margin direction"),
    ("Gross NPA", "Asset-quality risk"),
    ("Net NPA", "Credit-loss risk"),
    ("CASA Ratio", "Low-cost funding strength"),
    ("Capital Adequacy / CET1", "Balance-sheet resilience"),
    ("Management Guidance", "Forward-looking qualitative signal"),
]


def is_screener_workbook(path: Path) -> bool:
    try:
        return REQUIRED_SHEETS <= set(pd.ExcelFile(path).sheet_names)
    except Exception:  # noqa: BLE001 - not a readable workbook
        return False


def _f(v: Any) -> float | None:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    return None if math.isnan(x) or math.isinf(x) else x


def _col(df: pd.DataFrame, name: str) -> list[float]:
    if name not in df.columns:
        return []
    return [x for x in (_f(v) for v in df[name]) if x is not None]


def _cagr(vals: list[float], years: int = 3) -> float | None:
    vals = vals[-(years + 1):]
    if len(vals) < 2 or vals[0] <= 0 or vals[-1] <= 0:
        return None
    return (vals[-1] / vals[0]) ** (1 / (len(vals) - 1)) - 1


def _clip(x: float | None, lo: float, hi: float, default: float) -> float:
    return round(min(hi, max(lo, x if x is not None else default)), 3)


def _status(rule: str, x: float | None) -> str:
    if x is None:
        return "n/a"
    if rule == "growth_rev":
        return "Strong" if x >= 0.08 else "Stable" if x >= 0 else "Weak"
    if rule == "growth_profit":
        return "Strong" if x >= 0.10 else "Stable" if x >= 0 else "Weak"
    if rule == "margin":
        return "Strong" if x >= 0.42 else "Stable" if x >= 0.30 else "Watch"
    if rule == "roe":
        return "Strong" if x >= 0.15 else "Stable" if x >= 0.12 else "Watch"
    if rule == "pe":
        return "Attractive/Low" if x <= 15 else "Normal" if x <= 20 else "Rich"
    if rule == "cash":
        return "Positive" if x > 0 else "Weak"
    return "n/a"


def _overall(statuses: list[str]) -> str:
    bad = sum(s in ("Weak", "Watch") for s in statuses)
    good = sum(s in ("Strong", "Positive", "Attractive/Low") for s in statuses)
    if bad >= 3:
        return "Watch"
    if good >= 4:
        return "Positive"
    return "Neutral-Positive" if bad == 0 else "Neutral"


def build_prediction_model(path: Path) -> dict[str, Any]:
    """Raises ``ValueError`` when the workbook lacks the data needed."""
    pl = transpose_statement_sheet(read_statement_sheet_raw(path, "Profit & Loss"))
    q = transpose_statement_sheet(read_statement_sheet_raw(path, "Quarters"))
    ds = parse_label_value_block(read_statement_sheet_raw(path, "Data Sheet"), 0, 1)
    sheets = set(pd.ExcelFile(path).sheet_names)
    bs = (transpose_statement_sheet(read_statement_sheet_raw(path, "Balance Sheet"))
          if "Balance Sheet" in sheets else pd.DataFrame())
    cf = (transpose_statement_sheet(read_statement_sheet_raw(path, "Cash Flow"))
          if "Cash Flow" in sheets else pd.DataFrame())

    sales, npf, eps = _col(pl, "Sales"), _col(pl, "Net profit"), _col(pl, "EPS")
    if len(sales) < 2 or len(npf) < 2:
        raise ValueError("Profit & Loss sheet needs at least 2 periods of Sales and Net profit")
    qs, qn, qo = _col(q, "Sales"), _col(q, "Net profit"), _col(q, "Operating Profit")
    last_period = str(pl.index[-1])
    ttm_label, label_fn = fiscal_year_labels(last_period)
    fy = ttm_label.split(" ")[0]

    if len(qs) >= 4 and len(qn) >= 4:
        ttm_sales, ttm_np = sum(qs[-4:]), sum(qn[-4:])
        ttm_op = sum(qo[-4:]) if len(qo) >= 4 else None
        ttm_src = "Sum of last 4 quarters"
    else:
        ttm_sales, ttm_np, ttm_op = sales[-1], npf[-1], (_col(pl, "Operating Profit") or [None])[-1]
        ttm_src = "Latest fiscal year"
    shares = npf[-1] / eps[-1] if eps and eps[-1] else None
    if not shares:
        raise ValueError("EPS missing — cannot derive share count")
    ttm_eps = ttm_np / shares
    price = _f(ds.get("Current Price")) or (_col(pl, "Price") or [None])[-1]
    if not price:
        raise ValueError("Current Price missing from Data Sheet")
    pe = price / ttm_eps if ttm_eps > 0 else None

    equity = _col(bs, "Equity Share Capital")
    reserves = _col(bs, "Reserves")
    roe = npf[-1] / (equity[-1] + reserves[-1]) if equity and reserves else None
    ocf = (_col(cf, "Cash from Operating Activity") or [None])[-1]
    q_yoy_s = qs[-1] / qs[-5] - 1 if len(qs) >= 5 and qs[-5] else None
    q_yoy_n = qn[-1] / qn[-5] - 1 if len(qn) >= 5 and qn[-5] else None
    q_opm = qo[-1] / qs[-1] if qo and qs and qs[-1] else None
    ttm_margin = ttm_op / ttm_sales if ttm_op is not None and ttm_sales else None
    fy_sg = sales[-1] / sales[-2] - 1 if sales[-2] else None
    fy_ng = npf[-1] / npf[-2] - 1 if npf[-2] else None

    base_sg = _clip(_cagr(sales), 0.02, 0.20, 0.05)
    base_pg = _clip(_cagr(npf), 0.03, 0.25, 0.09)
    hist_pe = pe or 15.0
    sg = {"bear": round(base_sg * 0.5, 3), "base": base_sg, "bull": round(base_sg * 1.5, 3)}
    pg = {"bear": round(base_pg * 0.5, 3), "base": base_pg, "bull": round(base_pg * 1.5, 3)}
    tpe = {"bear": round(hist_pe * 0.85, 1), "base": round(hist_pe, 1),
           "bull": round(hist_pe * 1.2, 1)}

    scen = scenario_projection(ttm_sales, ttm_np, shares, price, sg, pg, tpe)
    outlook = multi_year_outlook(ttm_sales, ttm_np, shares, sg["base"], pg["base"], 3,
                                 start_label=ttm_label, year_label_fn=label_fn)

    signals = [
        ("Revenue momentum", _status("growth_rev", q_yoy_s), "Latest quarter sales YoY"),
        ("Profit momentum", _status("growth_profit", q_yoy_n), "Latest quarter net profit YoY"),
        ("Margin", _status("margin", ttm_margin), "TTM operating-profit margin"),
        ("ROE trend", _status("roe", roe), f"{fy} ROE level"),
        ("Valuation", _status("pe", pe), "Relative earnings multiple"),
        ("Cash generation", _status("cash", ocf), f"{fy} operating cash flow direction"),
    ]
    overall = _overall([s[1] for s in signals])
    snapshot = [
        (f"{fy} Sales Growth", fy_sg, "pct"), (f"{fy} Net Profit Growth", fy_ng, "pct"),
        ("Latest Quarter Sales YoY", q_yoy_s, "pct"),
        ("Latest Quarter Net Profit YoY", q_yoy_n, "pct"),
        ("TTM Operating Margin", ttm_margin, "pct"), (f"{fy} ROE", roe, "pct"),
        ("Current P/E", pe, "num"), ("Current Price", price, "num"),
    ]
    inputs = [
        (f"{fy} Sales", sales[-1], f"Profit & Loss {last_period}"),
        ("TTM Sales", ttm_sales, ttm_src),
        (f"{fy} Net Profit", npf[-1], f"Profit & Loss {last_period}"),
        ("TTM Net Profit", ttm_np, ttm_src),
        ("TTM EPS", ttm_eps, "TTM net profit / shares"),
        ("Current P/E", pe, "Price / TTM EPS"),
        ("Sheet Current Price", price, "Data Sheet"),
        ("Latest Quarter Sales", qs[-1] if qs else None, f"Quarters {q.index[-1]}"),
        ("Latest Quarter Net Profit", qn[-1] if qn else None, f"Quarters {q.index[-1]}"),
        ("Latest Quarter OPM", q_opm, "Quarters"),
        (f"{fy} ROE", roe, "Net profit / (equity + reserves)"),
        (f"{fy} Operating Cash Flow", ocf, "Cash Flow"),
    ]
    history = []
    for period, sale, profit in zip(pl.index, pl["Sales"], pl["Net profit"], strict=True):
        if _f(sale) is not None and _f(profit) is not None:
            history.append({"period": f"FY{str(period)[2:4]}", "sales": _f(sale),
                            "net_profit": _f(profit)})
    return {
        "history": history[-6:],
        "company": str(ds.get("COMPANY NAME") or path.stem), "fy": fy, "ttm_label": ttm_label,
        "inputs": inputs, "snapshot": snapshot, "signals": signals, "overall": overall,
        "assumptions": {"sales_growth": sg, "profit_growth": pg, "target_pe": tpe},
        "scenarios": scen, "outlook": outlook,
        "confidence": {"score": 50, "band": "Medium",
                       "note": "Base score; add the company-specific inputs below to raise it."},
        "shares": shares,
        "assumption_note": ("Defaults derived from this company's own history: base growth = "
                            "3-year CAGR (clipped); bear = 50%, bull = 150% of base; target P/E "
                            "= current P/E x 0.85 / 1.0 / 1.2. Edit them in the workbook."),
    }


# ─────────────────────────── HTML ───────────────────────────
def _n(v: float | None, pct: bool = False) -> str:
    if v is None:
        return "—"
    return f"{v * 100:.1f}%" if pct else f"{v:,.2f}"


def html_tables(m: dict[str, Any]) -> list[dict[str, Any]]:
    sc = m["scenarios"]
    a = m["assumptions"]
    tables: list[dict[str, Any]] = [
        {"heading": "1. Current snapshot & momentum", "note": "", "num": True,
         "columns": ["Metric", "Result"],
         "rows": [[lab, _n(v, kind == "pct")] for lab, v, kind in m["snapshot"]]},
        {"heading": "Signals", "num": False, "note": f"Overall historical signal: {m['overall']}",
         "columns": ["Signal", "Status", "Basis"],
         "rows": [[s, st, b] for s, st, b in m["signals"]]},
        {"heading": f"2. {m['outlook'][1]['period'].rstrip('E')} scenario forecast", "note": m["assumption_note"], "num": True,
         "columns": ["Metric", "Bear", "Base", "Bull"],
         "rows": [
             ["Sales growth", *[_n(a["sales_growth"][k], True) for k in ("bear", "base", "bull")]],
             ["Profit growth", *[_n(a["profit_growth"][k], True) for k in ("bear", "base", "bull")]],
             ["Sales", *[_n(sc[k]["sales"]) for k in ("bear", "base", "bull")]],
             ["Net profit", *[_n(sc[k]["net_profit"]) for k in ("bear", "base", "bull")]],
             ["EPS", *[_n(sc[k]["eps"]) for k in ("bear", "base", "bull")]],
             ["Target P/E", *[_n(sc[k]["target_pe"]) for k in ("bear", "base", "bull")]],
             ["Implied valuation price", *[_n(sc[k]["implied_price"]) for k in ("bear", "base", "bull")]],
             ["Upside / downside vs price",
              *[_n((sc[k]["upside_downside_pct"] or 0) / 100, True) for k in ("bear", "base", "bull")]],
         ]},
        {"heading": "3. Base-case multi-year outlook", "note": "", "num": True,
         "columns": ["Year", "Sales", "Net profit", "EPS"],
         "rows": [[o["period"], _n(o["sales"]), _n(o["net_profit"]), _n(o["eps"])]
                  for o in m["outlook"]]},
        {"heading": "Model confidence", "num": True, "note": m["confidence"]["note"],
         "columns": ["Score", "Band"],
         "rows": [[str(m["confidence"]["score"]), m["confidence"]["band"]]]},
    ]
    return tables


# ─────────────────────────── charts (HTML) ───────────────────────────
def charts_b64(m: dict[str, Any]) -> list[dict[str, str]]:
    """PNG graphs for the HTML report: history + base-case projection, and scenario prices."""
    import matplotlib.pyplot as plt

    from app.services.reports.charts import _fig_to_b64

    hist = m["history"]
    proj = m["outlook"][1:]
    labels = [h["period"] for h in hist] + [o["period"] for o in proj]
    out: list[dict[str, str]] = []

    fig, axes = plt.subplots(1, 2, figsize=(10, 3.8))
    for ax, key, title in ((axes[0], "sales", "Sales"), (axes[1], "net_profit", "Net profit")):
        actual = [h[key] for h in hist]
        projected = [o[key] for o in proj]
        ax.bar(range(len(actual)), actual, color="#2563eb", label="actual")
        ax.bar(range(len(actual), len(actual) + len(projected)), projected, color="#d97706",
               hatch="//", alpha=0.85, label="base-case projection")
        ax.set_xticks(range(len(labels)))
        ax.set_xticklabels(labels, fontsize=8)
        ax.set_title(f"{title}: history and projection", fontsize=10)
        ax.grid(axis="y", alpha=0.25)
        ax.legend(fontsize=8, loc="upper left")
    fig.tight_layout()
    out.append({"heading": "History and base-case projection", "b64": _fig_to_b64(fig)})

    sc = m["scenarios"]
    price = next(v for k, v, _ in m["inputs"] if k == "Sheet Current Price")
    names = ["bear", "base", "bull"]
    vals = [sc[k]["implied_price"] for k in names]
    fig, ax = plt.subplots(figsize=(6.5, 3.6))
    bars = ax.bar([n.title() for n in names], vals, color=["#b42318", "#6b7280", "#1a7f37"])
    ax.axhline(price, color="black", linestyle="--", linewidth=1,
               label=f"current price {price:,.0f}")
    for b, v in zip(bars, vals, strict=True):
        ax.text(b.get_x() + b.get_width() / 2, v, f"{v:,.0f}", ha="center", va="bottom",
                fontsize=9)
    ax.set_title(f"{m['outlook'][1]['period'].rstrip('E')} implied valuation price by scenario",
                 fontsize=10)
    ax.grid(axis="y", alpha=0.25)
    ax.legend(fontsize=8)
    fig.tight_layout()
    out.append({"heading": "Scenario valuation", "b64": _fig_to_b64(fig)})
    return out


# ─────────────────────────── Excel ───────────────────────────
def write_prediction_sheets(wb, m: dict[str, Any]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    head = PatternFill("solid", fgColor="1F3864")
    sub = PatternFill("solid", fgColor="D9E2F3")
    bold = Font(bold=True)

    def hdr(ws, row, cols):
        for c in cols:
            cell = ws.cell(row, c)
            cell.font, cell.fill = bold, sub

    co = m["company"]
    pi = wb.create_sheet("Prediction Inputs")
    pi["A1"] = f"{co} – FORECAST MODEL INPUTS"
    pi["A1"].font, pi["A1"].fill = Font(bold=True, size=13, color="FFFFFF"), head
    pi.merge_cells("A1:F1")
    for c, t in zip("ABCEF", ["Historical / Current Metric", "Value", "Source",
                              "Scenario Assumption", "Value"], strict=True):
        pi[f"{c}3"] = t
    hdr(pi, 3, [1, 2, 3, 5, 6])
    for i, (lab, val, src) in enumerate(m["inputs"]):
        r = 4 + i
        pi[f"A{r}"], pi[f"B{r}"], pi[f"C{r}"] = lab, val, src
        pi[f"B{r}"].number_format = "0.0%" if "OPM" in lab or "ROE" in lab else "#,##0.00"
    a = m["assumptions"]
    rows = [("Bear Sales Growth", a["sales_growth"]["bear"]), ("Base Sales Growth", a["sales_growth"]["base"]),
            ("Bull Sales Growth", a["sales_growth"]["bull"]), ("Bear Profit Growth", a["profit_growth"]["bear"]),
            ("Base Profit Growth", a["profit_growth"]["base"]), ("Bull Profit Growth", a["profit_growth"]["bull"]),
            ("Bear Target P/E", a["target_pe"]["bear"]), ("Base Target P/E", a["target_pe"]["base"]),
            ("Bull Target P/E", a["target_pe"]["bull"])]
    for i, (lab, val) in enumerate(rows):
        r = 4 + i
        pi[f"E{r}"], pi[f"F{r}"] = lab, val
        pi[f"F{r}"].number_format = "0.0%" if "Growth" in lab else "0.0"
        pi[f"F{r}"].fill = PatternFill("solid", fgColor="FFF2CC")
    pi["E14"] = m["assumption_note"]
    pi["E14"].alignment = Alignment(wrap_text=True, vertical="top")
    pi.merge_cells("E14:F18")

    pi["A18"] = "COMPANY-SPECIFIC INPUTS TO ADD FOR A STRONGER MODEL"
    pi["A18"].font = bold
    for c, t in zip("ABC", ["Parameter", "Current Value", "Why It Matters"], strict=True):
        pi[f"{c}20"] = t
    hdr(pi, 20, [1, 2, 3])
    pi["E20"], pi["F20"] = "Model Control", "Value"
    hdr(pi, 20, [5, 6])
    for i, (p, why) in enumerate(EXTERNAL_PARAMS):
        pi[f"A{21 + i}"], pi[f"C{21 + i}"] = p, why
        pi[f"B{21 + i}"].fill = PatternFill("solid", fgColor="FFF2CC")
    pi["E21"], pi["F21"] = "External inputs completed", "=COUNT(B21:B30)"
    pi["E22"], pi["F22"] = "Model confidence score", "=MIN(90,50+F21*4)"
    pi["E23"], pi["F23"] = "Confidence band", (
        '=IF(F22>=80,"High",IF(F22>=65,"Medium-High",IF(F22>=50,"Medium","Low")))')
    pi["E24"], pi["F24"] = "Forecast horizon", f"{m['outlook'][1]['period']}–{m['outlook'][-1]['period']}"
    pi["E25"], pi["F25"] = "Method", "Scenario + trend + momentum"
    for col, w in zip("ABCDEF", [30, 16, 36, 3, 30, 16], strict=True):
        pi.column_dimensions[col].width = w

    pr = wb.create_sheet("Prediction Report")
    pr["A1"] = f"{co} – DATA-DRIVEN PREDICTION REPORT"
    pr["A1"].font, pr["A1"].fill = Font(bold=True, size=13, color="FFFFFF"), head
    pr.merge_cells("A1:G1")
    pr["A2"] = "Built from the uploaded historical workbook. Forecasts are scenarios, not guaranteed outcomes."
    pr["A4"] = "1. CURRENT SNAPSHOT & MOMENTUM"
    pr["A4"].font = bold
    for c, t in zip("ABDE", ["Metric", "Result", "Signal", "Status"], strict=True):
        pr[f"{c}5"] = t
    hdr(pr, 5, [1, 2, 4, 5])
    for i, (lab, val, kind) in enumerate(m["snapshot"]):
        r = 6 + i
        pr[f"A{r}"], pr[f"B{r}"] = lab, val
        pr[f"B{r}"].number_format = "0.0%" if kind == "pct" else "#,##0.00"
    for i, (sig, st, basis) in enumerate(m["signals"]):
        pr[f"D{6 + i}"], pr[f"E{6 + i}"], pr[f"F{6 + i}"] = sig, st, basis
    pr["D12"] = "Overall historical signal"
    pr["E12"] = ('=IF(COUNTIF(E6:E11,"Weak")+COUNTIF(E6:E11,"Watch")>=3,"Watch",'
                 'IF(COUNTIF(E6:E11,"Strong")+COUNTIF(E6:E11,"Positive")+COUNTIF(E6:E11,"Attractive/Low")>=4,'
                 '"Positive",IF(COUNTIF(E6:E11,"Weak")+COUNTIF(E6:E11,"Watch")=0,"Neutral-Positive","Neutral")))')
    pr["F12"] = "Combined rule-based reading of the workbook"
    pr["D13"], pr["E13"] = "Model confidence", "='Prediction Inputs'!F22"

    pr["A16"] = f"2. {m['outlook'][1]['period'].rstrip('E')} SCENARIO FORECAST"
    pr["A16"].font = bold
    for c, t in zip("ABCD", ["Metric", "Bear Case", "Base Case", "Bull Case"], strict=True):
        pr[f"{c}18"] = t
    hdr(pr, 18, [1, 2, 3, 4])
    sg, pg, pe = ("F4", "F5", "F6"), ("F7", "F8", "F9"), ("F10", "F11", "F12")
    for j, col in enumerate("BCD"):
        pr[f"{col}19"] = f"='Prediction Inputs'!B5*(1+'Prediction Inputs'!{sg[j]})"
        pr[f"{col}20"] = f"='Prediction Inputs'!B7*(1+'Prediction Inputs'!{pg[j]})"
        pr[f"{col}21"] = f"='Prediction Inputs'!B8*(1+'Prediction Inputs'!{pg[j]})"
        pr[f"{col}22"] = f"='Prediction Inputs'!{pe[j]}"
        pr[f"{col}23"] = f"={col}21*{col}22"
        pr[f"{col}24"] = f"={col}23/'Prediction Inputs'!B10-1"
        for r in (19, 20, 21, 22, 23):
            pr[f"{col}{r}"].number_format = "#,##0.00"
        pr[f"{col}24"].number_format = "0.0%"
    for r, lab in zip(range(19, 25), ["Sales", "Net Profit", "EPS", "Target P/E",
                                      "Implied Valuation Price", "Upside / Downside vs Sheet Price"],
                      strict=True):
        pr[f"A{r}"] = lab
    pr["A25"] = "Sheet Current Price"
    for col in "BCD":
        pr[f"{col}25"] = "='Prediction Inputs'!B10"
        pr[f"{col}25"].number_format = "#,##0.00"

    pr["A27"] = "3. BASE-CASE MULTI-YEAR OUTLOOK"
    pr["A27"].font = bold
    for c, t in zip("ABCD", ["Year", "Sales", "Net Profit", "EPS"], strict=True):
        pr[f"{c}29"] = t
    hdr(pr, 29, [1, 2, 3, 4])
    pr["A30"], pr["B30"], pr["C30"], pr["D30"] = (
        m["outlook"][0]["period"], "='Prediction Inputs'!B5", "='Prediction Inputs'!B7",
        "='Prediction Inputs'!B8")
    for i, o in enumerate(m["outlook"][1:]):
        r = 31 + i
        pr[f"A{r}"] = o["period"]
        pr[f"B{r}"] = f"=B{r - 1}*(1+'Prediction Inputs'!F5)"
        pr[f"C{r}"] = f"=C{r - 1}*(1+'Prediction Inputs'!F8)"
        pr[f"D{r}"] = f"=D{r - 1}*(1+'Prediction Inputs'!F8)"
    for r in range(30, 30 + len(m["outlook"])):
        for col in "BCD":
            pr[f"{col}{r}"].number_format = "#,##0.00"
    pr["A36"] = ("This workbook alone is not enough for a high-confidence forecast. Add the "
                 "company-specific inputs on the Prediction Inputs sheet; do not treat an implied "
                 "valuation price as a guaranteed stock-price prediction.")
    pr["A36"].alignment = Alignment(wrap_text=True, vertical="top")
    pr.merge_cells("A36:G38")
    for col, w in zip("ABCDEFG", [34, 16, 16, 30, 18, 42, 8], strict=True):
        pr.column_dimensions[col].width = w

    _write_excel_charts(pr, m)


def _write_excel_charts(pr, m: dict[str, Any]) -> None:
    """Chart data block (history + formulas pointing at the outlook) and native charts."""
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font

    pr["A40"] = "4. CHART DATA (history + base-case projection; projection cells follow section 3)"
    pr["A40"].font = Font(bold=True)
    for c, t in zip("ABC", ["Year", "Sales", "Net Profit"], strict=True):
        pr[f"{c}41"] = t
        pr[f"{c}41"].font = Font(bold=True)
    row = 42
    for h in m["history"]:
        pr[f"A{row}"], pr[f"B{row}"], pr[f"C{row}"] = h["period"], h["sales"], h["net_profit"]
        row += 1
    for i in range(1, len(m["outlook"])):
        pr[f"A{row}"] = m["outlook"][i]["period"]
        pr[f"B{row}"] = f"=B{30 + i}"
        pr[f"C{row}"] = f"=C{30 + i}"
        row += 1
    last = row - 1
    for r in range(42, last + 1):
        pr[f"B{r}"].number_format = pr[f"C{r}"].number_format = "#,##0.00"

    cats = Reference(pr, min_col=1, min_row=42, max_row=last)
    for title, col, anchor in (("Sales: history and projection", 2, "I4"),
                               ("Net profit: history and projection", 3, "I22")):
        ch = BarChart()
        ch.type, ch.title = "col", title
        ch.height, ch.width = 8, 16
        ch.add_data(Reference(pr, min_col=col, min_row=41, max_row=last), titles_from_data=True)
        ch.set_categories(cats)
        ch.legend = None
        ch.series[0].graphicalProperties.solidFill = "2563EB"
        pr.add_chart(ch, anchor)

    sc = BarChart()
    sc.type, sc.title = "col", "Implied price by scenario vs current price"
    sc.height, sc.width = 8, 16
    sc.add_data(Reference(pr, min_col=1, max_col=4, min_row=23), from_rows=True,
                titles_from_data=True)
    sc.add_data(Reference(pr, min_col=1, max_col=4, min_row=25), from_rows=True,
                titles_from_data=True)
    sc.set_categories(Reference(pr, min_col=2, max_col=4, min_row=18))
    sc.series[0].graphicalProperties.solidFill = "1A7F37"
    sc.series[1].graphicalProperties.solidFill = "9CA3AF"
    pr.add_chart(sc, "I40")
