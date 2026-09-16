"""Write the calculation engine's own output into a workbook as a new sheet.

Complements ``app.services.reports.renderer`` (which renders the agent
pipeline's findings to HTML) — this writes the lower-level calculation
engine's numbers (scenario projection + statistics.full_report) as a sheet.

Caller's choice, and it matters: adding this sheet to a *fresh* ``Workbook()``
is always safe. Adding it to a workbook you loaded from an existing file and
plan to re-save is not, unless every sheet in it is data only — openpyxl
keeps a formula cell's formula but not its last-calculated value when a
workbook is loaded normally (not ``data_only=True``), so re-saving that
workbook drops every formula cell's cached display value elsewhere in the
file (the formulas themselves survive, so Excel/LibreOffice recalculates
them correctly on next open — but anything reading the file's raw values
before that, including our own parsers, would see blanks). See
``scripts/write_grd_calculation.py`` for the safe pattern: parse the source
workbook read-only, write this sheet into a brand new standalone workbook.
"""

from __future__ import annotations

from typing import Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

_TITLE_FONT = Font(bold=True, size=13)
_SECTION_FONT = Font(bold=True, size=11)
_HEADER_FONT = Font(bold=True)

_SCENARIO_ORDER = ("bear", "base", "bull")


def write_grd_calculation_sheet(
    wb: openpyxl.Workbook,
    *,
    scenario_rows: dict[str, dict[str, float]] | None = None,
    outlook_rows: list[dict[str, Any]] | None = None,
    data_driven: dict[str, Any] | None = None,
    ratio_report: dict[str, Any] | None = None,
    sheet_name: str = "GRD Calculation",
) -> Worksheet:
    """Create (or replace) ``sheet_name`` in ``wb`` with:

    1. the scenario projection (``scenario.scenario_projection`` output) —
       what-if next-period Sales/Net Profit/EPS/valuation per scenario
    2. the multi-year outlook (``scenario.multi_year_outlook`` output)
    3. the data-driven cross-check (``statistics.full_report`` output) — the
       historical trend and an ETS-based forecast, independent of the
       scenario assumptions above
    4. the ratio report (``ratios.full_ratio_report`` output) — Margins /
       Returns / Valuation / Quality

    Any section left as ``None`` is skipped. Does not save the workbook —
    call ``wb.save(path)`` after.
    """
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(sheet_name)

    row = 1
    ws.cell(row, 1, "GRD CALCULATION").font = _TITLE_FONT
    row += 1
    ws.cell(
        row, 1,
        "Computed by Grd-stk-mkt's calculation engine — "
        "a data-driven cross-check, not a guaranteed outcome.",
    )
    row += 2

    if scenario_rows:
        row = _write_scenario_section(ws, row, scenario_rows)
        row += 1

    if outlook_rows:
        row = _write_outlook_section(ws, row, outlook_rows)
        row += 1

    if data_driven:
        row = _write_data_driven_section(ws, row, data_driven)
        row += 1

    if ratio_report:
        row = _write_ratio_section(ws, row, ratio_report)

    for col_idx in range(1, 15):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 20
    return ws


def _write_scenario_section(
    ws: Worksheet, row: int, scenario_rows: dict[str, dict[str, float]]
) -> int:
    ws.cell(row, 1, "1. SCENARIO PROJECTION (next period)").font = _SECTION_FONT
    row += 1
    scenarios = [s for s in _SCENARIO_ORDER if s in scenario_rows] or list(scenario_rows)
    # matches the header wording on Screener's own "Prediction Report" sheet
    # ("Bear Case" / "Base Case" / "Bull Case"), not just "Bear" / "Base" / "Bull"
    headers = ["Metric", *[f"{s.capitalize()} Case" for s in scenarios]]
    for col, h in enumerate(headers, start=1):
        ws.cell(row, col, h).font = _HEADER_FONT
    row += 1

    metric_labels = [
        ("sales", "Sales"),
        ("net_profit", "Net Profit"),
        ("eps", "EPS"),
        ("target_pe", "Target P/E"),
        ("implied_price", "Implied Valuation Price"),
        ("upside_downside_pct", "Upside / Downside vs Current Price (%)"),
    ]
    for key, label in metric_labels:
        ws.cell(row, 1, label)
        for col, scen in enumerate(scenarios, start=2):
            ws.cell(row, col, scenario_rows[scen].get(key))
        row += 1
    return row


def _write_outlook_section(ws: Worksheet, row: int, outlook_rows: list[dict[str, Any]]) -> int:
    ws.cell(row, 1, "2. MULTI-YEAR OUTLOOK (base case, compounded)").font = _SECTION_FONT
    row += 1
    for col, h in enumerate(["Period", "Sales", "Net Profit", "EPS"], start=1):
        ws.cell(row, col, h).font = _HEADER_FONT
    row += 1
    for entry in outlook_rows:
        ws.cell(row, 1, entry["period"])
        ws.cell(row, 2, entry["sales"])
        ws.cell(row, 3, entry["net_profit"])
        ws.cell(row, 4, entry["eps"])
        row += 1
    return row


def _write_data_driven_section(ws: Worksheet, row: int, report: dict[str, Any]) -> int:
    ws.cell(
        row, 1,
        "3. DATA-DRIVEN CROSS-CHECK (historical trend, independent of scenario inputs above)",
    ).font = _SECTION_FONT
    row += 1
    headers = [
        "Metric", "Mean", "Median", "Std Dev", "Latest YoY %", "CAGR %",
        "Trend", "Forecast (next)", "95% CI Low", "95% CI High",
        "Max Drawdown %", "Downside Deviation",
        "Backtest MAPE %", "Backtest Hit Rate %",
    ]
    for col, h in enumerate(headers, start=1):
        ws.cell(row, col, h).font = _HEADER_FONT
    row += 1

    for metric, sections in report.get("metrics", {}).items():
        desc = sections.get("descriptive_stats", {})
        growth = sections.get("growth_trend", {})
        fc = sections.get("forecast", {})
        risk = sections.get("volatility_downside", {})
        bt = sections.get("backtest", {})
        forecast_point = fc.get("forecast", [None])[0] if not fc.get("insufficient_data") else None
        ci = fc.get("confidence_intervals", [{}])[0] if not fc.get("insufficient_data") else {}
        ws.cell(row, 1, metric)
        ws.cell(row, 2, desc.get("mean"))
        ws.cell(row, 3, desc.get("median"))
        ws.cell(row, 4, desc.get("std"))
        ws.cell(row, 5, growth.get("latest_yoy_pct"))
        ws.cell(row, 6, growth.get("cagr_pct"))
        ws.cell(row, 7, growth.get("trend_direction"))
        ws.cell(row, 8, forecast_point)
        ws.cell(row, 9, ci.get("low"))
        ws.cell(row, 10, ci.get("high"))
        ws.cell(row, 11, None if risk.get("insufficient_data") else risk.get("max_drawdown_pct"))
        ws.cell(row, 12, None if risk.get("insufficient_data") else risk.get("downside_deviation"))
        ws.cell(row, 13, None if bt.get("insufficient_data") else bt.get("mape_pct"))
        ws.cell(
            row, 14,
            None if bt.get("insufficient_data") else bt.get("directional_hit_rate_pct"),
        )
        row += 1

    if "regression" in report and not report["regression"].get("insufficient_data"):
        row += 1
        reg = report["regression"]
        ws.cell(row, 1, "Regression").font = _HEADER_FONT
        row += 1
        ws.cell(row, 1, f"{reg['target']} ~ {' + '.join(reg['features'])}")
        row += 1
        for feat, coef in reg.get("coefficients", {}).items():
            ws.cell(row, 1, f"  coefficient[{feat}]")
            ws.cell(row, 2, coef)
            row += 1
        ws.cell(row, 1, "  R-squared")
        ws.cell(row, 2, reg.get("r_squared"))
        row += 1

    if "correlation" in report:
        row += 1
        ws.cell(row, 1, "Correlation note").font = _HEADER_FONT
        row += 1
        ws.cell(row, 1, report["correlation"].get("note", ""))
        row += 1

    return row


_RATIO_LABELS: dict[str, list[tuple[str, str]]] = {
    "margins": [
        ("operating_profit_margin_pct", "Operating Profit Margin % (OPM)"),
        ("net_profit_margin_pct", "Net Profit Margin %"),
        ("net_interest_margin_pct", "Net Interest Margin % (NIM, banks)"),
    ],
    "returns": [
        ("roe_pct", "Return on Equity % (ROE)"),
        ("roa_pct", "Return on Assets % (ROA)"),
        ("roce_pct", "Return on Capital Employed % (ROCE)"),
    ],
    "valuation": [
        ("pe", "P/E"),
        ("pb", "P/B"),
        ("ev_to_ebitda", "EV / EBITDA"),
        ("ev_to_sales", "EV / Sales"),
    ],
    "quality": [
        ("cash_conversion", "Cash Conversion (CFO / Net Profit)"),
        ("leverage_debt_to_equity", "Leverage (Debt / Equity)"),
        ("asset_quality", "Asset Quality"),
    ],
}


def _write_ratio_section(ws: Worksheet, row: int, ratio_report: dict[str, Any]) -> int:
    ws.cell(
        row, 1,
        "4. RATIOS (Margins / Returns / Valuation / Quality)",
    ).font = _SECTION_FONT
    row += 1
    for col, h in enumerate(["Category", "Ratio", "Latest", "Note"], start=1):
        ws.cell(row, col, h).font = _HEADER_FONT
    row += 1

    for category, items in _RATIO_LABELS.items():
        section = ratio_report.get(category, {})
        for key, label in items:
            r = section.get(key, {})
            ws.cell(row, 1, category.capitalize())
            ws.cell(row, 2, label)
            if r.get("insufficient_data"):
                ws.cell(row, 3, None)
                ws.cell(row, 4, r.get("reason", "insufficient data"))
            else:
                ws.cell(row, 3, r.get("latest"))
                ws.cell(row, 4, r.get("note", ""))
            row += 1
    return row
