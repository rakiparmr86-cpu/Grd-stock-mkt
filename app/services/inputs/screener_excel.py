"""Transpose a Screener.in-style financial statement sheet into a tidy,
period-indexed DataFrame the statistics engine can consume.

Screener's export (Profit & Loss / Quarters / Balance Sheet / Cash Flow
sheets) is shaped one row per metric, one column per period — the opposite of
what ``app.services.calculations.statistics`` expects (one row per period,
one column per metric). This module does the transpose; it does not know
anything about tickers, storage, or the rest of the input-connector layer.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import openpyxl
import pandas as pd

# Screener appends forward-looking / scenario columns after the real periods —
# these aren't a period at all, so drop them rather than treat them as one.
_IGNORE_TRAILING_COLS = {"trailing", "best case", "worst case"}

# Every cell on a Screener statement sheet (Profit & Loss / Quarters /
# Balance Sheet / Cash Flow) is a bare pointer to one cell on the hidden
# "Data Sheet" tab, e.g. ``='Data Sheet'!B16`` — not a computed expression.
_SIMPLE_CELL_REF = re.compile(r"^='?([^'!]+?)'?!(\$?[A-Za-z]+\$?\d+)$")


def _resolve_simple_ref(formula: object, wb_values: openpyxl.Workbook) -> object:
    """Resolve a bare cross-sheet cell reference formula against a workbook
    already loaded with ``data_only=True``. Returns ``None`` for anything
    that isn't exactly that shape (arithmetic, functions, ranges) — those are
    genuinely out of scope; only Screener's own pointer-only cells are
    handled here."""
    if not isinstance(formula, str):
        return None
    m = _SIMPLE_CELL_REF.match(formula.strip())
    if not m:
        return None
    sheet_name, cell_ref = m.group(1), m.group(2).replace("$", "")
    if sheet_name not in wb_values.sheetnames:
        return None
    return wb_values[sheet_name][cell_ref].value


def read_statement_sheet_raw(path: Path, sheet: str) -> pd.DataFrame:
    """Like ``pd.read_excel(path, sheet_name=sheet, header=None)``, except a
    blank *cached* formula result is recovered from the cell it directly
    points to, elsewhere in the same workbook, when possible.

    Screener's statement sheets are built entirely from formulas referencing
    a hidden "Data Sheet" tab. If whatever last saved a particular copy of
    the file didn't let Excel recalculate first, every one of those cached
    results comes back blank — even though the underlying "Data Sheet" data
    they point to is still there and unaffected (it holds real values, not
    formulas). A plain ``pd.read_excel`` only ever sees the blank cache;
    this recovers what a person opening the file in Excel would see, without
    needing a full formula-evaluation engine — only exact single-cell
    pointers are resolved.
    """
    wb_values = openpyxl.load_workbook(path, data_only=True)
    wb_formulas = openpyxl.load_workbook(path, data_only=False)
    ws_values, ws_formulas = wb_values[sheet], wb_formulas[sheet]

    rows = []
    for r in range(1, ws_values.max_row + 1):
        row = []
        for c in range(1, ws_values.max_column + 1):
            value = ws_values.cell(r, c).value
            if value is None:
                value = _resolve_simple_ref(ws_formulas.cell(r, c).value, wb_values)
            row.append(value)
        rows.append(row)
    return pd.DataFrame(rows)


def transpose_statement_sheet(raw: pd.DataFrame) -> pd.DataFrame:
    """``raw`` is a sheet read with ``header=None`` (so row/column position is
    preserved). Finds the row whose first cell is "Narration" — that's the
    period header — and builds one row per period, one column per metric.
    """
    header_row_idx = None
    for i in range(len(raw)):
        first_cell = raw.iat[i, 0]
        if isinstance(first_cell, str) and first_cell.strip().lower() == "narration":
            header_row_idx = i
            break
    if header_row_idx is None:
        raise ValueError("could not find a 'Narration' header row in this sheet")

    header = raw.iloc[header_row_idx]
    period_cols: list[int] = []
    for col_idx in range(1, len(header)):
        label = header.iloc[col_idx]
        if label is None or (isinstance(label, float) and pd.isna(label)):
            continue
        if isinstance(label, str) and label.strip().lower() in _IGNORE_TRAILING_COLS:
            continue
        period_cols.append(col_idx)
    if not period_cols:
        raise ValueError("no period columns found after the 'Narration' header")

    periods: list[str] = []
    for col_idx in period_cols:
        label = header.iloc[col_idx]
        if isinstance(label, dt.datetime):
            periods.append(label.strftime("%Y-%m"))
        elif isinstance(label, float) and label.is_integer():
            periods.append(str(int(label)))  # a bare year (2019) read back as 2019.0
        else:
            periods.append(str(label).strip())

    metrics: dict[str, list[float | None]] = {}
    for i in range(header_row_idx + 1, len(raw)):
        name = raw.iat[i, 0]
        if name is None or (isinstance(name, float) and pd.isna(name)):
            continue
        name = str(name).strip()
        if not name:
            continue
        values: list[float | None] = []
        for col_idx in period_cols:
            v = raw.iat[i, col_idx]
            is_number = isinstance(v, (int, float)) and not isinstance(v, bool)
            values.append(float(v) if is_number else None)
        # a repeated metric name (rare, but Screener sheets aren't guaranteed
        # unique) would silently overwrite — keep the first occurrence
        metrics.setdefault(name, values)

    return pd.DataFrame(metrics, index=periods)


def parse_label_value_block(raw: pd.DataFrame, label_col: int, value_col: int) -> dict[str, object]:
    """Read a simple two-column "label, value" block (as used by the
    "Prediction Inputs" sheet's assumption tables) into a dict, skipping
    blank/non-string labels. Later duplicate labels overwrite earlier ones —
    harmless for a hand-maintained sheet with one row per named input."""
    out: dict[str, object] = {}
    n_cols = raw.shape[1]
    if label_col >= n_cols:
        return out
    for i in range(len(raw)):
        label = raw.iat[i, label_col]
        if not isinstance(label, str):
            continue
        label = label.strip()
        if not label:
            continue
        value = raw.iat[i, value_col] if value_col < n_cols else None
        out[label] = None if isinstance(value, float) and pd.isna(value) else value
    return out


def load_prediction_inputs(raw: pd.DataFrame) -> dict[str, object]:
    """Parse a "Prediction Inputs" sheet shaped like:
    ``Historical / Current Metric | Value | Source | | Scenario Assumption | Value``
    — two independent label/value tables side by side (columns 0-1 and 4-5).

    Returns the specific fields ``scenario_projection``/``multi_year_outlook``
    need; raises ``ValueError`` naming any that are missing so a malformed or
    differently-labeled sheet fails clearly instead of silently.
    """
    actuals = parse_label_value_block(raw, 0, 1)
    scenarios = parse_label_value_block(raw, 4, 5)

    def _require(block: dict[str, object], key: str) -> float:
        if key not in block or block[key] is None:
            raise ValueError(f"'Prediction Inputs' sheet is missing required value: {key!r}")
        return float(block[key])  # type: ignore[arg-type]

    ttm_net_profit = _require(actuals, "TTM Net Profit")
    ttm_eps = _require(actuals, "TTM EPS")
    return {
        "ttm_sales": _require(actuals, "TTM Sales"),
        "ttm_net_profit": ttm_net_profit,
        "current_price": _require(actuals, "Sheet Current Price"),
        "shares_outstanding": ttm_net_profit / ttm_eps,
        "sales_growth": {
            "bear": _require(scenarios, "Bear Sales Growth"),
            "base": _require(scenarios, "Base Sales Growth"),
            "bull": _require(scenarios, "Bull Sales Growth"),
        },
        "profit_growth": {
            "bear": _require(scenarios, "Bear Profit Growth"),
            "base": _require(scenarios, "Base Profit Growth"),
            "bull": _require(scenarios, "Bull Profit Growth"),
        },
        "target_pe": {
            "bear": _require(scenarios, "Bear Target P/E"),
            "base": _require(scenarios, "Base Target P/E"),
            "bull": _require(scenarios, "Bull Target P/E"),
        },
    }
