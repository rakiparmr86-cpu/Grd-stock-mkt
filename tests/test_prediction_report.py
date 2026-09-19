"""Prediction Inputs / Prediction Report built from a Screener-style workbook."""

from __future__ import annotations

import datetime as dt
import io

import openpyxl
import pytest

from app.services.reports.prediction_report import (
    build_prediction_model,
    html_tables,
    is_screener_workbook,
    write_prediction_sheets,
)


def _workbook(path, *, price=100.0):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    pl = wb.create_sheet("Profit & Loss")
    pl.append(["ACME LTD"])
    pl.append(["Narration", *[dt.datetime(y, 3, 31) for y in (2022, 2023, 2024, 2025, 2026)]])
    pl.append(["Sales", 100, 110, 121, 133, 146])
    pl.append(["Operating Profit", 30, 33, 36, 40, 44])
    pl.append(["Net profit", 10, 11, 12.1, 13.3, 14.6])
    pl.append(["EPS", 1.0, 1.1, 1.21, 1.33, 1.46])
    q = wb.create_sheet("Quarters")
    q.append(["ACME LTD"])
    q.append(["Narration", *[dt.datetime(2025, m, 28) for m in (3, 6, 9, 12)],
              *[dt.datetime(2026, m, 28) for m in (3, 6)]])
    q.append(["Sales", 30, 32, 34, 36, 38, 40])
    q.append(["Operating Profit", 9, 10, 10, 11, 12, 13])
    q.append(["Net profit", 3, 3.2, 3.4, 3.6, 3.8, 4.0])
    ds = wb.create_sheet("Data Sheet")
    ds.append(["COMPANY NAME", "ACME LTD"])
    ds.append(["Current Price", price])
    wb.save(path)
    return path


def test_detects_screener_workbook(tmp_path):
    assert is_screener_workbook(_workbook(tmp_path / "a.xlsx"))
    other = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    wb.save(other)
    assert not is_screener_workbook(other)


def test_model_scenarios_are_ordered_and_consistent(tmp_path):
    m = build_prediction_model(_workbook(tmp_path / "a.xlsx"))
    assert m["company"] == "ACME LTD"
    sc = m["scenarios"]
    assert sc["bear"]["implied_price"] < sc["base"]["implied_price"] < sc["bull"]["implied_price"]
    ttm_sales = 34 + 36 + 38 + 40
    assert m["inputs"][1][1] == ttm_sales
    assert m["outlook"][0]["period"] == "FY26 / TTM"
    assert m["outlook"][1]["period"] == "FY27E"
    assert m["overall"] in {"Watch", "Positive", "Neutral-Positive", "Neutral"}


def test_missing_price_raises(tmp_path):
    p = _workbook(tmp_path / "a.xlsx", price=None)
    with pytest.raises(ValueError, match="Current Price"):
        build_prediction_model(p)


def test_html_tables_and_excel_layout(tmp_path):
    m = build_prediction_model(_workbook(tmp_path / "a.xlsx"))
    headings = [t["heading"] for t in html_tables(m)]
    assert any("scenario forecast" in h for h in headings)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_prediction_sheets(wb, m)
    buf = io.BytesIO()
    wb.save(buf)
    wb2 = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
    assert wb2.sheetnames == ["Prediction Inputs", "Prediction Report"]
    pr = wb2["Prediction Report"]
    assert pr["A18"].value == "Metric" and pr["B18"].value == "Bear Case"
    assert pr["B23"].value == "=B21*B22"
    assert wb2["Prediction Inputs"]["A5"].value == "TTM Sales"
