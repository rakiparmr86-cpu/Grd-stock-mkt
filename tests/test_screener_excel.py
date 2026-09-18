from __future__ import annotations

import datetime as dt

import openpyxl
import pandas as pd
import pytest

from app.services.inputs.screener_excel import (
    load_prediction_inputs,
    read_statement_sheet_raw,
    transpose_statement_sheet,
)


def _sheet(rows: list[list]) -> pd.DataFrame:
    """Build a raw (header=None) sheet DataFrame from row lists, padded like
    ``pd.read_excel`` would (ragged rows become NaN)."""
    return pd.DataFrame(rows)


def test_transpose_annual_sheet():
    raw = _sheet([
        ["HDFC BANK LTD"],
        [
            "Narration", dt.datetime(2022, 3, 1), dt.datetime(2023, 3, 1),
            "Trailing", "Best Case", "Worst Case",
        ],
        ["Sales", 100.0, 121.0, 130.0, 140.0, 120.0],
        ["Net profit", 10.0, 15.0, 16.0, 18.0, 12.0],
    ])
    df = transpose_statement_sheet(raw)
    assert list(df.index) == ["2022-03", "2023-03"]
    assert list(df.columns) == ["Sales", "Net profit"]
    assert df.loc["2022-03", "Sales"] == 100.0
    assert df.loc["2023-03", "Net profit"] == 15.0


def test_transpose_ignores_trailing_scenario_columns():
    raw = _sheet([
        ["Narration", "Q1", "Q2", "Trailing", "Best Case", "Worst Case"],
        ["Sales", 1.0, 2.0, 3.0, 4.0, 5.0],
    ])
    df = transpose_statement_sheet(raw)
    assert list(df.columns) == ["Sales"]
    assert list(df.index) == ["Q1", "Q2"]
    assert df["Sales"].tolist() == [1.0, 2.0]


def test_transpose_bare_numeric_year_header_not_dotzero():
    # openpyxl/pandas read a plain numeric-looking header like "2019" back
    # as the float 2019.0 rather than a string — don't leak that as "2019.0"
    raw = _sheet([
        ["Narration", 2019.0, 2020.0],
        ["Sales", 100.0, 110.0],
    ])
    df = transpose_statement_sheet(raw)
    assert list(df.index) == ["2019", "2020"]


def test_transpose_missing_narration_row_raises():
    raw = _sheet([["not a header", 1, 2]])
    with pytest.raises(ValueError, match="Narration"):
        transpose_statement_sheet(raw)


def test_transpose_skips_blank_metric_rows():
    raw = _sheet([
        ["Narration", "2022", "2023"],
        [None, None, None],
        ["Sales", 100.0, 110.0],
    ])
    df = transpose_statement_sheet(raw)
    assert list(df.columns) == ["Sales"]


def test_transpose_non_numeric_cell_becomes_none():
    raw = _sheet([
        ["Narration", "2022", "2023"],
        ["EPS", 10.5, "NM"],
    ])
    df = transpose_statement_sheet(raw)
    assert df.loc["2022", "EPS"] == 10.5
    assert pd.isna(df.loc["2023", "EPS"])


# ── Prediction Inputs sheet ────────────────────────────────────────────────
def _prediction_inputs_sheet() -> pd.DataFrame:
    return _sheet([
        ["HDFC BANK - FORECAST MODEL INPUTS"],
        [None],
        ["Historical / Current Metric", "Value", "Source", None, "Scenario Assumption", "Value"],
        ["TTM Sales", 351818.61, "Profit & Loss Trailing", None, "Bear Sales Growth", 0.01],
        ["TTM Net Profit", 79012.77, "Profit & Loss Trailing", None, "Base Sales Growth", 0.05],
        ["TTM EPS", 51.27457397295704, "Profit & Loss Trailing", None, "Bull Sales Growth", 0.08],
        ["Sheet Current Price", 708, "Data Sheet", None, "Bear Profit Growth", 0.03],
        [None, None, None, None, "Base Profit Growth", 0.09],
        [None, None, None, None, "Bull Profit Growth", 0.14],
        [None, None, None, None, "Bear Target P/E", 12.5],
        [None, None, None, None, "Base Target P/E", 15],
        [None, None, None, None, "Bull Target P/E", 17],
    ])


def test_load_prediction_inputs_parses_both_blocks():
    parsed = load_prediction_inputs(_prediction_inputs_sheet())
    assert parsed["ttm_sales"] == pytest.approx(351818.61)
    assert parsed["ttm_net_profit"] == pytest.approx(79012.77)
    assert parsed["current_price"] == pytest.approx(708)
    assert parsed["shares_outstanding"] == pytest.approx(79012.77 / 51.27457397295704)
    assert parsed["sales_growth"] == {"bear": 0.01, "base": 0.05, "bull": 0.08}
    assert parsed["profit_growth"] == {"bear": 0.03, "base": 0.09, "bull": 0.14}
    assert parsed["target_pe"] == {"bear": 12.5, "base": 15.0, "bull": 17.0}


def test_load_prediction_inputs_missing_field_raises():
    raw = _sheet([["Historical / Current Metric", "Value"], ["TTM Sales", 100.0]])
    with pytest.raises(ValueError, match="TTM Net Profit"):
        load_prediction_inputs(raw)


# ── read_statement_sheet_raw: stale/uncalculated formula caches ────────────
def _workbook_with_stale_formula_caches(path) -> None:
    """A minimal reproduction of a real bug found live: a Screener export
    whose statement-sheet cells are all bare pointers to a hidden "Data
    Sheet" tab (``='Data Sheet'!B16``), saved without Excel ever
    recalculating them — so every cached value openpyxl/pandas would
    normally read back is blank, even though "Data Sheet" itself holds real,
    non-formula values."""
    wb = openpyxl.Workbook()
    data_sheet = wb.active
    data_sheet.title = "Data Sheet"
    # row 1: period headers; row 2: Sales values — arbitrary layout, mirrors
    # the real template's "one row per line item, referenced by row number" shape
    data_sheet.append(["", dt.datetime(2022, 3, 31), dt.datetime(2023, 3, 31)])
    data_sheet.append(["", 100.0, 121.0])

    ws = wb.create_sheet("Profit & Loss")
    ws.append(["DEMO CO LTD"])
    ws.append(["Narration", "='Data Sheet'!B1", "='Data Sheet'!C1"])
    ws.append(["Sales", "='Data Sheet'!B2", "='Data Sheet'!C2"])
    wb.save(path)


def test_read_statement_sheet_raw_resolves_stale_formula_caches(tmp_path):
    path = tmp_path / "stale.xlsx"
    _workbook_with_stale_formula_caches(path)

    # sanity check: a plain read sees nothing (this is the bug being fixed)
    plain = pd.read_excel(path, sheet_name="Profit & Loss", header=None)
    with pytest.raises(ValueError, match="no period columns"):
        transpose_statement_sheet(plain)

    raw = read_statement_sheet_raw(path, "Profit & Loss")
    wide = transpose_statement_sheet(raw)
    assert list(wide.index) == ["2022-03", "2023-03"]
    assert wide.loc["2022-03", "Sales"] == 100.0
    assert wide.loc["2023-03", "Sales"] == 121.0


def test_read_statement_sheet_raw_leaves_non_pointer_formulas_alone(tmp_path):
    """A formula more complex than a bare cell pointer (arithmetic, a
    function call) is genuinely out of scope — it stays blank rather than
    being (incorrectly) evaluated."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Label", "Value"])
    ws.append(["Total", "=1+1"])
    path = tmp_path / "formula.xlsx"
    wb.save(path)

    raw = read_statement_sheet_raw(path, "Sheet1")
    assert pd.isna(raw.iat[1, 1])
