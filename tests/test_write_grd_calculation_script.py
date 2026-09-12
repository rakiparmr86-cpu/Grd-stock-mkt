"""End-to-end test for scripts/write_grd_calculation.py — the piece that
matters most here is that the SOURCE workbook is never touched, since an
earlier version of this script mutated it in place via openpyxl, which
silently drops every formula cell's cached value on save (see
app/services/reports/excel_writeback.py's module docstring)."""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import openpyxl
import pytest

_SCRIPT_PATH = Path(__file__).resolve().parents[1] / "scripts" / "write_grd_calculation.py"
_spec = importlib.util.spec_from_file_location("write_grd_calculation", _SCRIPT_PATH)
write_grd_calculation = importlib.util.module_from_spec(_spec)
sys.modules[_spec.name] = write_grd_calculation
_spec.loader.exec_module(write_grd_calculation)


def _build_source_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    pl = wb.active
    pl.title = "Profit & Loss"
    pl.append(["DEMO CO"])
    pl.append(["Narration", 2019, 2020, 2021, 2022, 2023, "Trailing", "Best Case", "Worst Case"])
    pl.append(["Sales", 100, 110, 121, 133, 146, 150, 160, 140])
    # a formula cell, like Screener's real sheets pull from "Data Sheet"
    pl.append(["Net profit", 10, "=B3*0.12", 15, 17, 20, 21, 23, 18])

    inputs = wb.create_sheet("Prediction Inputs")
    inputs.append([
        "Historical / Current Metric", "Value", "Source", None, "Scenario Assumption", "Value",
    ])
    inputs.append(["TTM Sales", 150.0, "x", None, "Bear Sales Growth", 0.01])
    inputs.append(["TTM Net Profit", 21.0, "x", None, "Base Sales Growth", 0.05])
    inputs.append(["TTM EPS", 2.1, "x", None, "Bull Sales Growth", 0.08])
    inputs.append(["Sheet Current Price", 50.0, "x", None, "Bear Profit Growth", 0.03])
    inputs.append([None, None, None, None, "Base Profit Growth", 0.09])
    inputs.append([None, None, None, None, "Bull Profit Growth", 0.14])
    inputs.append([None, None, None, None, "Bear Target P/E", 12.5])
    inputs.append([None, None, None, None, "Base Target P/E", 15])
    inputs.append([None, None, None, None, "Bull Target P/E", 17])
    wb.save(path)


def test_script_never_modifies_source_file(tmp_path):
    source = tmp_path / "source.xlsx"
    _build_source_workbook(source)
    original_bytes = source.read_bytes()
    original_mtime = source.stat().st_mtime_ns

    output = tmp_path / "out.xlsx"
    rc = write_grd_calculation.main([str(source), "--output", str(output)])

    assert rc == 0
    assert source.read_bytes() == original_bytes
    assert source.stat().st_mtime_ns == original_mtime


def test_script_output_has_grd_calculation_sheet_only(tmp_path):
    source = tmp_path / "source.xlsx"
    _build_source_workbook(source)
    output = tmp_path / "out.xlsx"

    write_grd_calculation.main([str(source), "--output", str(output)])

    wb = openpyxl.load_workbook(output)
    assert wb.sheetnames == ["GRD Calculation"]
    ws = wb["GRD Calculation"]
    flat_col1 = [c.value for row in ws.iter_rows() for c in [row[0]]]
    assert "1. SCENARIO PROJECTION (next period)" in flat_col1
    assert "2. MULTI-YEAR OUTLOOK (base case, compounded)" in flat_col1


def test_script_default_output_path(tmp_path):
    source = tmp_path / "source.xlsx"
    _build_source_workbook(source)

    write_grd_calculation.main([str(source)])

    expected = tmp_path / "source.GRD_Calculation.xlsx"
    assert expected.exists()


def test_script_errors_on_missing_source(tmp_path, capsys):
    rc = write_grd_calculation.main([str(tmp_path / "nope.xlsx")])
    assert rc == 1
    assert "does not exist" in capsys.readouterr().err


def test_script_warns_and_skips_data_driven_on_missing_target(tmp_path, capsys):
    source = tmp_path / "source.xlsx"
    _build_source_workbook(source)
    output = tmp_path / "out.xlsx"

    rc = write_grd_calculation.main([
        str(source), "--output", str(output), "--target", "Nonexistent Metric",
    ])

    assert rc == 0
    assert "skipping data-driven section" in capsys.readouterr().err
    wb = openpyxl.load_workbook(output)
    flat_col1 = [c.value for row in wb["GRD Calculation"].iter_rows() for c in [row[0]]]
    assert not any("DATA-DRIVEN" in str(v) for v in flat_col1 if v)


@pytest.mark.parametrize("field", ["TTM Sales", "TTM Net Profit", "TTM EPS", "Sheet Current Price"])
def test_script_errors_clearly_on_malformed_inputs_sheet(tmp_path, field):
    source = tmp_path / "source.xlsx"
    _build_source_workbook(source)
    wb = openpyxl.load_workbook(source)
    ws = wb["Prediction Inputs"]
    for row in ws.iter_rows():
        if row[0].value == field:
            row[1].value = None
    wb.save(source)

    with pytest.raises(ValueError, match=field):
        write_grd_calculation.main([str(source)])
