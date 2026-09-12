"""ExcelConnector's fundamental row_kind: flat tidy rows (existing behavior)
and auto-detected Screener-style row-per-metric statement sheets (new)."""

from __future__ import annotations

import openpyxl

from app.services.inputs.base import ConnectorKind
from app.services.inputs.excel import ExcelConnector


def test_flat_fundamental_rows_unchanged(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ticker", "period", "revenue", "net_income"])
    ws.append(["ACME", "FY2024", 100.0, 10.0])
    path = tmp_path / "flat.xlsx"
    wb.save(path)

    conn = ExcelConnector({"path": str(path), "mode": "rows", "row_kind": "fundamental"})
    results = list(conn.fetch())
    assert len(results) == 1
    rows = results[0].rows
    assert rows.iloc[0]["ticker"] == "ACME"
    assert rows.iloc[0]["revenue"] == 100.0


def test_flat_rows_get_ticker_from_config_when_missing(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["period", "revenue"])
    ws.append(["FY2024", 100.0])
    path = tmp_path / "flat_no_ticker.xlsx"
    wb.save(path)

    conn = ExcelConnector({
        "path": str(path), "mode": "rows", "row_kind": "fundamental", "ticker": "ACME",
    })
    rows = next(iter(conn.fetch())).rows
    assert rows.iloc[0]["ticker"] == "ACME"


def _screener_style_workbook(path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Profit & Loss"
    ws.append(["DEMO CO LTD"])
    ws.append([
        "Narration", "2021", "2022", "2023", "2024",
        "Trailing", "Best Case", "Worst Case",
    ])
    ws.append(["Sales", 100, 110, 121, 133, 140, 150, 130])
    ws.append(["Net profit", 10, 12, 15, 17, 18, 20, 15])
    ws.append(["EPS", 1.0, 1.2, 1.5, 1.7, 1.8, 2.0, 1.5])
    wb.save(path)


def test_screener_shape_auto_detected_and_transposed(tmp_path):
    path = tmp_path / "screener.xlsx"
    _screener_style_workbook(path)

    conn = ExcelConnector({
        "path": str(path), "mode": "rows", "row_kind": "fundamental",
        "sheet": "Profit & Loss", "ticker": "DEMO",
    })
    results = list(conn.fetch())
    assert len(results) == 1
    rows = results[0].rows
    assert set(rows["period"]) == {"2021", "2022", "2023", "2024"}
    assert (rows["ticker"] == "DEMO").all()

    row_2024 = rows[rows["period"] == "2024"].iloc[0]
    assert row_2024["revenue"] == 133.0   # "Sales" -> revenue
    assert row_2024["net_income"] == 17.0  # "Net profit" -> net_income
    assert row_2024["eps"] == 1.7


def test_screener_shape_requires_ticker(tmp_path):
    path = tmp_path / "screener.xlsx"
    _screener_style_workbook(path)

    # no ticker in config -> falls back to flat-row parsing of the same
    # sheet, which won't look anything like the transposed shape
    conn = ExcelConnector({
        "path": str(path), "mode": "rows", "row_kind": "fundamental", "sheet": "Profit & Loss",
    })
    rows = next(iter(conn.fetch())).rows
    assert "revenue" not in rows.columns


def test_default_kind_is_rows_for_fundamental_mode(tmp_path):
    path = tmp_path / "screener.xlsx"
    _screener_style_workbook(path)
    conn = ExcelConnector({
        "path": str(path), "mode": "rows", "row_kind": "fundamental",
        "sheet": "Profit & Loss", "ticker": "DEMO",
    })
    assert conn.default_kind == ConnectorKind.ROWS
