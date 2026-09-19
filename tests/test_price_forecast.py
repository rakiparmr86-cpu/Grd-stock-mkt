"""Price forecast service + its exports. The frame loader is stubbed so no DB is needed."""

from __future__ import annotations

import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from app.services import price_forecast as pf


def _frame(n=120):
    idx = pd.date_range("2026-01-01", periods=n, freq="B")
    close = 100 + np.arange(n) * 0.5 + np.sin(np.arange(n) / 5)
    return pd.DataFrame({"open": close - 0.3, "high": close + 1, "low": close - 1,
                         "close": close, "volume": 1000.0}, index=idx)


@pytest.fixture
def stub(monkeypatch):
    monkeypatch.setattr(pf, "load_ohlcv_frame", lambda ticker, limit=750: _frame().tail(limit))


def test_forecast_shape_and_dates(stub):
    fc = pf.build_price_forecast("acme", history=60, ahead=10)
    assert fc["ticker"] == "ACME"
    assert len(fc["bars"]) == 60 and len(fc["forecast"]) == 10
    assert fc["forecast"][0]["ts"] > fc["bars"][-1]["ts"]
    assert all(p["low"] <= p["value"] <= p["high"] for p in fc["forecast"])
    assert fc["forecast_end"] > fc["last_close"]  # rising series extrapolates upward


def test_too_few_bars_is_a_clear_error(monkeypatch):
    monkeypatch.setattr(pf, "load_ohlcv_frame", lambda t, limit=750: _frame(10))
    with pytest.raises(pf.ForecastError, match="at least 30"):
        pf.build_price_forecast("acme")


def test_no_data_is_a_clear_error(monkeypatch):
    monkeypatch.setattr(pf, "load_ohlcv_frame", lambda t, limit=750: pd.DataFrame())
    with pytest.raises(pf.ForecastError, match="no price data"):
        pf.build_price_forecast("acme")


def test_html_has_chart_and_tables(stub):
    html = pf.forecast_html(pf.build_price_forecast("acme", history=60, ahead=10))
    assert "data:image/png;base64" in html
    assert "Forecast path" in html and "ACME" in html


def test_excel_has_data_and_native_chart(stub):
    data = pf.forecast_excel(pf.build_price_forecast("acme", history=60, ahead=10))
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Forecast", "Summary"]
    ws = wb["Forecast"]
    assert ws["A4"].value == "Date" and ws["D4"].value == "Forecast"
    assert ws.max_row == 4 + 60 + 10
    assert len(ws._charts) == 1
