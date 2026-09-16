from __future__ import annotations

import openpyxl
import pandas as pd
import pytest

from app.services.calculations.ratios import full_ratio_report
from app.services.calculations.scenario import multi_year_outlook, scenario_projection
from app.services.calculations.statistics import full_report
from app.services.reports.excel_writeback import write_grd_calculation_sheet


@pytest.fixture
def scenario_and_outlook():
    scenario_rows = scenario_projection(
        ttm_sales=100.0, ttm_net_profit=10.0, shares_outstanding=5.0, current_price=20.0,
        sales_growth={"bear": 0.01, "base": 0.05, "bull": 0.08},
        profit_growth={"bear": 0.03, "base": 0.09, "bull": 0.14},
        target_pe={"bear": 12.5, "base": 15, "bull": 17},
    )
    outlook_rows = multi_year_outlook(100.0, 10.0, 5.0, 0.05, 0.09, years=3)
    return scenario_rows, outlook_rows


def test_write_sheet_creates_expected_structure(tmp_path, scenario_and_outlook):
    scenario_rows, outlook_rows = scenario_and_outlook
    wb = openpyxl.Workbook()
    write_grd_calculation_sheet(wb, scenario_rows=scenario_rows, outlook_rows=outlook_rows)

    assert "GRD Calculation" in wb.sheetnames
    ws = wb["GRD Calculation"]
    values = [[c.value for c in row] for row in ws.iter_rows()]
    flat_col1 = [row[0] for row in values]
    assert "GRD CALCULATION" in flat_col1
    assert "1. SCENARIO PROJECTION (next period)" in flat_col1
    assert "2. MULTI-YEAR OUTLOOK (base case, compounded)" in flat_col1
    assert "Sales" in flat_col1

    out_path = tmp_path / "out.xlsx"
    wb.save(out_path)
    reloaded = openpyxl.load_workbook(out_path)
    assert "GRD Calculation" in reloaded.sheetnames


def test_scenario_section_header_matches_screeners_own_wording(scenario_and_outlook):
    """The scenario table's column headers must read "Bear Case" / "Base
    Case" / "Bull Case" — the exact wording Screener's own "Prediction
    Report" sheet uses (section "2. FY27 SCENARIO FORECAST") — not just
    "Bear" / "Base" / "Bull", so a reader can't tell the two apart."""
    scenario_rows, _ = scenario_and_outlook
    wb = openpyxl.Workbook()
    write_grd_calculation_sheet(wb, scenario_rows=scenario_rows)
    ws = wb["GRD Calculation"]

    header_row = next(
        [c.value for c in row] for row in ws.iter_rows() if row[0].value == "Metric"
    )
    assert header_row[:4] == ["Metric", "Bear Case", "Base Case", "Bull Case"]


def test_write_sheet_replaces_existing_sheet_of_same_name(scenario_and_outlook):
    scenario_rows, outlook_rows = scenario_and_outlook
    wb = openpyxl.Workbook()
    write_grd_calculation_sheet(wb, scenario_rows=scenario_rows, outlook_rows=outlook_rows)
    first_count = len(wb.sheetnames)
    write_grd_calculation_sheet(wb, scenario_rows=scenario_rows, outlook_rows=outlook_rows)
    assert len(wb.sheetnames) == first_count  # replaced, not duplicated


def test_write_sheet_preserves_other_sheets(scenario_and_outlook):
    scenario_rows, outlook_rows = scenario_and_outlook
    wb = openpyxl.Workbook()
    wb.active.title = "Profit & Loss"
    wb["Profit & Loss"]["A1"] = "keep me"
    write_grd_calculation_sheet(wb, scenario_rows=scenario_rows, outlook_rows=outlook_rows)
    assert wb["Profit & Loss"]["A1"].value == "keep me"
    assert "GRD Calculation" in wb.sheetnames


def test_write_sheet_includes_data_driven_section():
    years = [str(y) for y in range(2019, 2024)]
    df = pd.DataFrame({
        "sales": [100.0, 110.0, 121.0, 133.0, 146.0],
        "net_profit": [10.0, 12.0, 15.0, 17.0, 20.0],
    }, index=years)
    report = full_report(df, target="net_profit", features=["sales"], periods_per_year=1)

    wb = openpyxl.Workbook()
    write_grd_calculation_sheet(wb, data_driven=report)
    ws = wb["GRD Calculation"]
    flat_col1 = [c.value for row in ws.iter_rows() for c in [row[0]]]
    assert any("DATA-DRIVEN CROSS-CHECK" in str(v) for v in flat_col1 if v)
    assert "sales" in flat_col1
    assert "net_profit" in flat_col1


def test_write_sheet_includes_risk_columns_from_volatility_downside():
    """The data-driven section must surface volatility_downside's fields
    (max drawdown, downside deviation) — full_report already computes them,
    but they used to be silently dropped instead of written to the sheet."""
    years = [str(y) for y in range(2019, 2024)]
    # monotonically increasing except a dip in year 3, so max_drawdown_pct is
    # non-zero and predictable, and downside_deviation is non-zero too.
    df = pd.DataFrame({
        "sales": [100.0, 120.0, 90.0, 130.0, 150.0],
        "net_profit": [10.0, 12.0, 15.0, 17.0, 20.0],
    }, index=years)
    report = full_report(df, target="net_profit", features=["sales"], periods_per_year=1)
    expected_risk = report["metrics"]["sales"]["volatility_downside"]
    assert expected_risk["max_drawdown_pct"] < 0
    assert expected_risk["downside_deviation"] > 0

    wb = openpyxl.Workbook()
    write_grd_calculation_sheet(wb, data_driven=report)
    ws = wb["GRD Calculation"]

    header_row = next(
        [c.value for c in row] for row in ws.iter_rows() if row[0].value == "Metric"
    )
    assert "Max Drawdown %" in header_row
    assert "Downside Deviation" in header_row
    dd_col = header_row.index("Max Drawdown %") + 1
    dev_col = header_row.index("Downside Deviation") + 1

    sales_row = next(row for row in ws.iter_rows() if row[0].value == "sales")
    assert sales_row[dd_col - 1].value == pytest.approx(expected_risk["max_drawdown_pct"])
    assert sales_row[dev_col - 1].value == pytest.approx(expected_risk["downside_deviation"])


def test_write_sheet_includes_ratio_section():
    years = ["2021", "2022", "2023"]
    df = pd.DataFrame({
        "revenue": [1000.0, 1100.0, 1250.0],
        "net_income": [100.0, 120.0, 150.0],
        "operating profit": [180.0, 200.0, 230.0],
        "equity share capital": [50.0, 50.0, 50.0],
        "reserves": [450.0, 550.0, 680.0],
        "total assets": [1200.0, 1350.0, 1500.0],
        "borrowings": [300.0, 280.0, 260.0],
    }, index=years)
    ratio_report = full_ratio_report(df, price=180.0)

    wb = openpyxl.Workbook()
    write_grd_calculation_sheet(wb, ratio_report=ratio_report)
    ws = wb["GRD Calculation"]
    flat_col1 = [c.value for row in ws.iter_rows() for c in [row[0]]]
    assert any("RATIOS" in str(v) for v in flat_col1 if v)

    header_row = next(
        [c.value for c in row] for row in ws.iter_rows() if row[0].value == "Category"
    )
    assert header_row[:4] == ["Category", "Ratio", "Latest", "Note"]

    roe_row = next(
        [c.value for c in row] for row in ws.iter_rows()
        if row[1].value == "Return on Equity % (ROE)"
    )
    expected_roe = ratio_report["returns"]["roe_pct"]["latest"]
    assert roe_row[2] == pytest.approx(expected_roe)
